import os
from pypdf import PdfReader
from docx import Document
from fastapi import HTTPException
from app.utils.text_cleaner import clean_text

def extract_text_from_pdf(file_path: str) -> str:
    """Extract and normalize text content from a PDF document."""
    try:
        reader = PdfReader(file_path)
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)
        
        full_text = "\n\n".join(extracted_pages)
        cleaned = clean_text(full_text)
        if not cleaned or len(cleaned.strip()) < 20:
            raise HTTPException(
                status_code=422,
                detail="PDF appears to be empty or contains only scanned images without selectable text. Please upload a searchable PDF or DOCX file."
            )
        return cleaned
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Failed to read PDF document: {str(e)}"
        )

def extract_text_from_docx(file_path: str) -> str:
    """Extract and normalize text content from a DOCX document."""
    try:
        doc = Document(file_path)
        paragraphs = []
        
        # Extract body paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        full_text = "\n".join(paragraphs)
        cleaned = clean_text(full_text)
        if not cleaned or len(cleaned.strip()) < 20:
            raise HTTPException(
                status_code=422,
                detail="DOCX file appears to be empty or has no extractable text."
            )
        return cleaned
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Failed to read Word (.docx) document: {str(e)}"
        )

def parse_document(file_path: str, ext: str) -> str:
    """Dispatches document extraction based on extension."""
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext == "docx":
        return extract_text_from_docx(file_path)
    else:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format: {ext}"
        )
